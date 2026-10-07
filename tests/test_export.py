"""Tests for CSV export (Phase 15) and rate limiting (Phase 17)."""
from datetime import date

from app.models.models import Account, Category, TransactionType
from app.services.finance import create_transaction

from tests.conftest import client, default_user_id, get_test_db  # noqa: E402,F401


def _add_expense(amount=35000, description="Nasi goreng"):
    db = get_test_db()
    acc = db.query(Account).filter(Account.name == "BCA").first()
    cat = db.query(Category).filter(Category.name == "Makan & Minum").first()
    create_transaction(db, user_id=default_user_id(), type=TransactionType.EXPENSE,
                       amount=amount, account_id=acc.id, category_id=cat.id,
                       date_val=date.today(), description=description)
    db.close()


def test_export_transactions_csv_scoped_to_user():
    _add_expense(35000)
    r = client.get("/transactions/export")
    assert r.status_code == 200
    assert "text/csv" in r.headers["content-type"]
    assert "transaksi_" in r.headers["content-disposition"]
    body = r.text
    assert "Tanggal" in body           # header row
    # BOM present so Excel opens UTF-8 correctly
    assert body.startswith("\ufeff")
    assert "Nasi goreng" in body
    assert "Makan & Minum" in body
    assert "Rp 35.000" in body


def test_export_transactions_empty():
    r = client.get("/transactions/export")
    assert r.status_code == 200
    assert "Tanggal" in r.text
    # no data rows beyond the header
    assert len([l for l in r.text.strip().split("\n") if l]) == 1


def test_export_reports_csv():
    _add_expense(35000)
    r = client.get("/reports/export")
    assert r.status_code == 200
    assert "text/csv" in r.headers["content-type"]
    assert r.text.startswith("\ufeff")
    assert "Laporan Keuangan" in r.text
    assert "Makan & Minum" in r.text
    assert "Rp 35.000" in r.text


def _read_xlsx(payload: bytes):
    """Parse an .xlsx response body into a list of rows (openpyxl)."""
    import io

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(payload), read_only=True)
    return [list(row) for row in wb.active.iter_rows(values_only=True)]


def test_export_transactions_xlsx():
    _add_expense(35000)
    r = client.get("/transactions/export.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "transaksi_" in r.headers["content-disposition"]
    assert r.headers["content-disposition"].endswith('.xlsx"')
    rows = _read_xlsx(r.content)
    assert rows[0][:6] == ["Tanggal", "Tipe", "Deskripsi", "Kategori", "Akun", "Jumlah"]
    flat = [c for row in rows for c in row if c]
    assert any("Nasi goreng" in str(c) for c in flat)
    assert any("Makan & Minum" in str(c) for c in flat)
    assert any("Rp 35.000" in str(c) for c in flat)


def test_export_transactions_xlsx_empty():
    r = client.get("/transactions/export.xlsx")
    assert r.status_code == 200
    rows = _read_xlsx(r.content)
    assert len(rows) == 1  # header only


def test_export_reports_xlsx():
    _add_expense(35000)
    r = client.get("/reports/export.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    rows = _read_xlsx(r.content)
    flat = [str(c) for row in rows for c in row if c]
    assert any("Laporan Keuangan" in c for c in flat)
    assert any("Makan & Minum" in c for c in flat)


def test_rate_limit_auth_blocks_after_limit():
    """Rapid repeated logins are throttled with HTTP 429."""
    # The shared client is already authenticated for other tests via the
    # dependency override; the rate limiter keys on client IP, so fire
    # actual /login POSTs (they will fail auth but still rate-limit).
    for _ in range(5):
        client.post("/login", data={
            "username": "bob", "password": "wrongpass",
            "csrf_token": "x", "next": "/",
        })
    r = client.post("/login", data={
        "username": "bob", "password": "wrongpass",
        "csrf_token": "x", "next": "/",
    })
    assert r.status_code == 429
    assert "Terlalu banyak" in r.text
    from app.rate_limit import RateLimitMiddleware
    RateLimitMiddleware.reset()


def test_rate_limit_does_not_block_normal_routes():
    for _ in range(3):
        r = client.get("/transactions/export")
        assert r.status_code == 200
    r = client.get("/")
    assert r.status_code == 200